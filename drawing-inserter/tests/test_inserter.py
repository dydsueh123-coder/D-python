import pytest
from pathlib import Path
from docx import Document
from core.inserter import insert_drawings, InsertConfig


def make_test_docx(tmp_path) -> Path:
    doc = Document()
    doc.add_paragraph("기존 본문 내용")
    p = tmp_path / "test.docx"
    doc.save(p)
    return p


def make_test_images(tmp_path, names: list[str]) -> list[Path]:
    from PIL import Image
    paths = []
    for name in names:
        img = Image.new("RGB", (100, 100), color=(200, 200, 200))
        p = tmp_path / name
        img.save(p)
        paths.append(p)
    return paths


def test_이미지_삽입_후_문서_저장(tmp_path):
    docx_path = make_test_docx(tmp_path)
    img_paths = make_test_images(tmp_path, ["a(1).png", "b(2).png"])
    output_path = tmp_path / "output.docx"

    config = InsertConfig()
    result = insert_drawings(docx_path, img_paths, output_path, config)

    assert output_path.exists()
    assert result["inserted"] == 2


def test_원본_문서_유지(tmp_path):
    docx_path = make_test_docx(tmp_path)
    img_paths = make_test_images(tmp_path, ["a(1).png"])
    output_path = tmp_path / "output.docx"

    original_mtime = docx_path.stat().st_mtime
    config = InsertConfig()
    insert_drawings(docx_path, img_paths, output_path, config)

    assert docx_path.stat().st_mtime == original_mtime


def test_캡션_텍스트_삽입_확인(tmp_path):
    docx_path = make_test_docx(tmp_path)
    img_paths = make_test_images(tmp_path, ["a(1).png"])
    output_path = tmp_path / "output.docx"

    config = InsertConfig()
    insert_drawings(docx_path, img_paths, output_path, config)

    doc = Document(output_path)
    texts = [p.text for p in doc.paragraphs]
    assert "<도면1>" in texts
