from dataclasses import dataclass
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
from core.parser import sort_drawings_with_failures
import shutil


@dataclass
class InsertConfig:
    max_width_cm: float = 14.5
    max_height_cm: float = 22.0
    caption_font: str = "맑은 고딕"
    caption_size_pt: int = 12
    caption_template: str = "【도 {key}】"


def _get_image_dimensions(img_path: Path, config: InsertConfig) -> tuple[float, float]:
    """원본 비율 유지하면서 최대 크기 안에 맞는 width, height (cm) 계산"""
    with Image.open(img_path) as img:
        orig_w, orig_h = img.size  # pixels

    # pixels → cm 변환 (96 dpi 기준, 1인치=2.54cm)
    dpi = 96
    w_cm = orig_w / dpi * 2.54
    h_cm = orig_h / dpi * 2.54

    if w_cm > config.max_width_cm:
        ratio = config.max_width_cm / w_cm
        w_cm *= ratio
        h_cm *= ratio

    if h_cm > config.max_height_cm:
        ratio = config.max_height_cm / h_cm
        w_cm *= ratio
        h_cm *= ratio

    return w_cm, h_cm


def insert_drawings(
    docx_path: Path,
    image_paths: list[Path],
    output_path: Path,
    config: InsertConfig,
    progress_callback=None,  # progress_callback(current, total) 형식
) -> dict:
    """
    원본 docx를 복사해서 이미지를 문서 끝에 삽입한다.
    원본 파일은 절대 수정하지 않는다.
    """
    # 원본 복사 (원본 보존)
    shutil.copy2(docx_path, output_path)
    doc = Document(output_path)

    filenames = [p.name for p in image_paths]
    path_map = {p.name: p for p in image_paths}

    sorted_drawings, failures = sort_drawings_with_failures(filenames)
    total = len(sorted_drawings)
    inserted = 0

    for i, drawing in enumerate(sorted_drawings):
        img_path = path_map[drawing["filename"]]
        key = drawing["key"]
        caption_text = config.caption_template.format(key=key)

        # 페이지 나누기 (각 도면을 새 페이지에 배치)
        page_break_para = doc.add_paragraph()
        page_break_run = page_break_para.add_run()
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        page_break_run._r.append(br)

        # 캡션 삽입 (왼쪽 정렬)
        caption_para = doc.add_paragraph()
        caption_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = caption_para.add_run(caption_text)
        run.font.name = config.caption_font
        run.font.size = Pt(config.caption_size_pt)

        # 이미지 삽입 (왼쪽 정렬)
        w_cm, _ = _get_image_dimensions(img_path, config)
        img_para = doc.add_paragraph()
        img_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        img_run = img_para.add_run()
        img_run.add_picture(str(img_path), width=Cm(w_cm))

        inserted += 1
        if progress_callback:
            progress_callback(i + 1, total)

    doc.save(output_path)

    return {
        "inserted": inserted,
        "failures": failures,
        "total_input": len(image_paths),
    }
